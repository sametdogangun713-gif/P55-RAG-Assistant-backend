"""Hafta 6 testleri: ayristirma, parcalama (chunking), yukleme is kurallari, sahiplik."""
import io
import os
import tempfile
import unittest
from pathlib import Path

from fastapi import HTTPException, UploadFile

from app.api import documents as docs_api
from app.core import config
from app.db import chunks as chunks_repo
from app.db import documents as documents_repo
from app.services import chunker, parser
from app.services import documents as svc
from tests.helpers import TempUploads, make_conn, make_user

try:
    from reportlab.pdfgen import canvas
except ImportError:
    canvas = None

try:
    import docx as _docx
except ImportError:
    _docx = None

UZUN_METIN = " ".join(
    f"Bu {i}. cumledir ve belge tabanli soru cevap sistemi icin ornek bir icerik tasir." for i in range(80)
)


class ChunkerTests(unittest.TestCase):
    def test_kisa_metin_tek_parca(self):
        parts = chunker.split_text("Merhaba dunya. Ikinci cumle.", 600, 100)
        self.assertEqual(parts, ["Merhaba dunya. Ikinci cumle."])

    def test_bos_metin_parca_uretmez(self):
        self.assertEqual(chunker.split_text("  \n\n \t ", 600, 100), [])
        self.assertEqual(chunker.chunk_pages([(1, ""), (2, "   ")]), [])

    def test_parcalar_siniri_asmaz_ve_icerik_kaybolmaz(self):
        parts = chunker.split_text(UZUN_METIN, 300, 60)
        self.assertGreater(len(parts), 3)
        self.assertTrue(all(len(p) <= 300 for p in parts))
        birlesik = " ".join(parts)
        for i in range(80):
            self.assertIn(f"Bu {i}. cumledir", birlesik)

    def test_ardisik_parcalar_ortusur(self):
        parts = chunker.split_text(UZUN_METIN, 300, 60)
        for a, b in zip(parts, parts[1:]):
            son_cumle = a.split(". ")[-1][:25]
            self.assertIn(son_cumle, b)

    def test_ortusme_sifirsa_tekrar_yok(self):
        parts = chunker.split_text(UZUN_METIN, 300, 0)
        self.assertEqual(" ".join(parts).count("Bu 5. cumledir"), 1)

    def test_noktalamasiz_dev_metin_zorla_bolunur(self):
        parts = chunker.split_text("kelime " * 2000, 200, 20)
        self.assertTrue(all(len(p) <= 200 for p in parts))
        parts = chunker.split_text("x" * 1000, 200, 20)       # bosluksuz tek kelime
        self.assertTrue(all(len(p) <= 200 for p in parts))
        self.assertEqual("".join(parts), "x" * 1000)

    def test_gecersiz_parametreler(self):
        with self.assertRaises(ValueError):
            chunker.split_text("abc", 600, 600)
        with self.assertRaises(ValueError):
            chunker.split_text("abc", 10, 0)

    def test_sayfa_no_ve_indeks_korunur(self):
        out = chunker.chunk_pages([(1, "Birinci sayfa."), (2, UZUN_METIN[:900])], 300, 50)
        self.assertEqual([c.index for c in out], list(range(len(out))))
        self.assertEqual(out[0].page_no, 1)
        self.assertTrue(all(c.page_no == 2 for c in out[1:]))

    def test_turkce_karakterler_korunur(self):
        out = chunker.split_text("Çalışkan öğrenci ğüşiöç İstanbul'a gitti.", 600, 100)
        self.assertIn("Çalışkan öğrenci ğüşiöç İstanbul'a", out[0])


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def _yaz(self, ad, data):
        p = Path(self.tmp) / ad
        p.write_bytes(data)
        return p

    def test_txt_utf8_ve_cp1254(self):
        metin = "Çalışma öğrencisi ğ ş ı İ"
        self.assertEqual(parser.parse_file(self._yaz("a.txt", metin.encode("utf-8")), ".txt")[0][1], metin)
        self.assertEqual(parser.parse_file(self._yaz("b.txt", metin.encode("cp1254")), ".txt")[0][1], metin)

    def test_desteklenmeyen_tur(self):
        with self.assertRaises(parser.ParseError):
            parser.parse_file(self._yaz("a.exe", b"x"), ".exe")

    def test_bozuk_pdf_ve_docx_ParseError(self):
        with self.assertRaises(parser.ParseError):
            parser.parse_file(self._yaz("k.pdf", b"%PDF-bozuk"), ".pdf")
        with self.assertRaises(parser.ParseError):
            parser.parse_file(self._yaz("k.docx", b"PK-bozuk"), ".docx")

    @unittest.skipIf(canvas is None, "reportlab kurulu degil")
    def test_pdf_sayfalar_ve_metin(self):
        p = Path(self.tmp) / "d.pdf"
        c = canvas.Canvas(str(p))
        c.drawString(72, 750, "Birinci sayfa metni")
        c.showPage()
        c.drawString(72, 750, "Ikinci sayfa metni")
        c.showPage()
        c.save()
        pages = parser.parse_file(p, ".pdf")
        self.assertEqual([n for n, _ in pages], [1, 2])
        self.assertIn("Birinci sayfa", pages[0][1])
        self.assertIn("Ikinci sayfa", pages[1][1])

    @unittest.skipIf(_docx is None, "python-docx kurulu degil")
    def test_docx_paragraf_ve_tablo(self):
        p = Path(self.tmp) / "d.docx"
        d = _docx.Document()
        d.add_paragraph("Paragraf bir")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Hucre A"
        t.rows[0].cells[1].text = "Hucre B"
        d.save(str(p))
        text = parser.parse_file(p, ".docx")[0][1]
        self.assertIn("Paragraf bir", text)
        self.assertIn("Hucre A | Hucre B", text)


class BelgeServisTests(unittest.TestCase):
    def setUp(self):
        self.conn = make_conn()
        self.u1 = make_user(self.conn, "bir@example.com")
        self.u2 = make_user(self.conn, "iki@example.com")
        self.admin = make_user(self.conn, "admin@example.com", role="admin")
        self._tmp = TempUploads()
        self.dir = self._tmp.__enter__()

    def tearDown(self):
        self._tmp.__exit__()

    def test_txt_yukleme_parcalanir_ve_diske_uuid_ile_yazilir(self):
        doc = svc.upload_document(self.conn, self.u1, "ders notu.txt", UZUN_METIN.encode("utf-8"), "text/plain")
        self.assertIn(doc["status"], ("chunked", "indexed"))   # Hafta 7'den itibaren parcalar vektorlestirilir -> 'indexed'
        self.assertGreater(doc["chunk_count"], 1)
        self.assertEqual(doc["filename"], "ders notu.txt")
        kayitli = documents_repo.get_document(self.conn, doc["id"])["stored_name"]
        self.assertTrue((Path(self.dir) / kayitli).is_file())
        self.assertNotIn("ders", kayitli)
        self.assertNotIn("stored_name", svc.public_document(doc))

    def test_dogrulama_kurallari(self):
        with self.assertRaises(svc.DocumentValidationError):
            svc.upload_document(self.conn, self.u1, "virus.exe", b"MZ...")
        with self.assertRaises(svc.DocumentValidationError):
            svc.upload_document(self.conn, self.u1, "bos.txt", b"")
        with self.assertRaises(svc.DocumentValidationError):
            svc.upload_document(self.conn, self.u1, "sahte.pdf", b"bu pdf degil")
        with self.assertRaises(svc.DocumentValidationError):
            svc.upload_document(self.conn, self.u1, "", b"x")
        self.assertEqual(documents_repo.list_documents(self.conn), [])   # reddedilenler kaydedilmez

    def test_boyut_siniri(self):
        eski = config.MAX_UPLOAD_MB
        config.MAX_UPLOAD_MB = 1
        try:
            svc.upload_document(self.conn, self.u1, "tam.txt", b"a" * (1024 * 1024))
            with self.assertRaises(svc.FileTooLargeError):
                svc.upload_document(self.conn, self.u1, "buyuk.txt", b"a" * (1024 * 1024 + 1))
        finally:
            config.MAX_UPLOAD_MB = eski

    def test_yol_gecisi_denemesi_etkisiz(self):
        doc = svc.upload_document(self.conn, self.u1, "../../etc/passwd.txt", b"icerik")
        self.assertEqual(doc["filename"], "passwd.txt")
        self.assertEqual(len(os.listdir(self.dir)), 1)
        self.assertEqual(svc.sanitize_filename("C:\\Users\\x\\not.txt"), "not.txt")

    @unittest.skipIf(canvas is None, "reportlab kurulu degil")
    def test_metinsiz_pdf_basarisiz_durumuna_duser(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.showPage()
        c.save()
        doc = svc.upload_document(self.conn, self.u1, "bos.pdf", buf.getvalue())
        self.assertEqual(doc["status"], "failed")
        self.assertIn("metin", doc["error"])

    def test_bozuk_docx_failed_olur_sunucu_hatasi_vermez(self):
        doc = svc.upload_document(self.conn, self.u1, "bozuk.docx", b"PK\x03\x04bozuk")
        self.assertEqual(doc["status"], "failed")

    def test_sahiplik_kontrolu(self):
        doc = svc.upload_document(self.conn, self.u1, "a.txt", b"Merhaba dunya.")
        with self.assertRaises(svc.DocumentNotFoundError):
            svc.get_owned_document(self.conn, self.u2, doc["id"])       # baskasi goremez
        self.assertEqual(svc.get_owned_document(self.conn, self.admin, doc["id"])["id"], doc["id"])
        with self.assertRaises(svc.DocumentNotFoundError):
            svc.get_owned_document(self.conn, self.u1, 9999)             # olmayan: ayni hata
        self.assertEqual(svc.list_user_documents(self.conn, self.u2), [])
        with self.assertRaises(svc.DocumentNotFoundError):
            svc.list_document_chunks(self.conn, self.u2, doc["id"])
        with self.assertRaises(svc.DocumentNotFoundError):
            svc.delete_document(self.conn, self.u2, doc["id"])

    def test_silme_dosya_ve_parcalari_kaldirir(self):
        doc = svc.upload_document(self.conn, self.u1, "a.txt", UZUN_METIN.encode())
        svc.delete_document(self.conn, self.u1, doc["id"])
        self.assertEqual(os.listdir(self.dir), [])
        self.assertEqual(chunks_repo.count_chunks(self.conn, doc["id"]), 0)
        self.assertIsNone(documents_repo.get_document(self.conn, doc["id"]))


class BelgeApiTests(unittest.TestCase):
    def setUp(self):
        self.conn = make_conn()
        self.u1 = make_user(self.conn, "bir@example.com")
        self.u2 = make_user(self.conn, "iki@example.com")
        self._tmp = TempUploads()
        self._tmp.__enter__()

    def tearDown(self):
        self._tmp.__exit__()

    def _dosya(self, ad, data):
        return UploadFile(file=io.BytesIO(data), filename=ad)

    def test_yukle_listele_getir_sil(self):
        out = docs_api.upload(self._dosya("a.txt", UZUN_METIN.encode()), self.u1, self.conn)
        self.assertIn(out["status"], ("chunked", "indexed"))
        self.assertEqual(len(docs_api.list_documents(self.u1, self.conn)), 1)
        self.assertEqual(docs_api.get_document(out["id"], self.u1, self.conn)["filename"], "a.txt")
        parcalar = docs_api.list_chunks(out["id"], 2, 0, self.u1, self.conn)
        self.assertEqual(len(parcalar), 2)
        self.assertEqual(docs_api.delete_document(out["id"], self.u1, self.conn), {"deleted": True})

    def test_hata_kodlari(self):
        with self.assertRaises(HTTPException) as e:
            docs_api.upload(self._dosya("x.exe", b"MZ"), self.u1, self.conn)
        self.assertEqual(e.exception.status_code, 400)
        out = docs_api.upload(self._dosya("a.txt", b"Merhaba."), self.u1, self.conn)
        with self.assertRaises(HTTPException) as e:
            docs_api.get_document(out["id"], self.u2, self.conn)
        self.assertEqual(e.exception.status_code, 404)
        eski = config.MAX_UPLOAD_MB
        config.MAX_UPLOAD_MB = 1
        try:
            with self.assertRaises(HTTPException) as e:
                docs_api.upload(self._dosya("b.txt", b"a" * (1024 * 1024 + 5)), self.u1, self.conn)
            self.assertEqual(e.exception.status_code, 413)
        finally:
            config.MAX_UPLOAD_MB = eski


if __name__ == "__main__":
    unittest.main()
