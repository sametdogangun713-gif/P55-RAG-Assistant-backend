"""Belge ayristirma: dosyadan duz metin cikarir. Sonuc: [(sayfa_no | None, metin), ...]"""
import zipfile
from pathlib import Path

from app.core import config


class ParseError(Exception):
    """Dosya okunamadi / bozuk / desteklenmiyor."""


def parse_file(path, ext: str) -> list:
    """[(sayfa_no | None, metin), ...] dondurur.

    NUL (\x00) karakterleri atilir: metin icin anlamsizdirlar ve PostgreSQL TEXT sutunu onlari kabul etmez
    (bulutta ikili/bozuk bir dosya yuklenince 500 hatasi veriyordu; Postgres testleri yakaladi).
    """
    path = Path(path)
    ext = ext.lower()
    if ext == ".txt":
        pages = _parse_txt(path)
    elif ext == ".pdf":
        pages = _parse_pdf(path)
    elif ext == ".docx":
        pages = _parse_docx(path)
    else:
        raise ParseError(f"Desteklenmeyen dosya türü: {ext}")
    return [(page, text.replace("\x00", "")) for page, text in pages]


def _parse_txt(path: Path) -> list:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1254"):   # cp1254: eski Turkce Windows kodlamasi
        try:
            return [(None, data.decode(enc))]
        except UnicodeDecodeError:
            continue
    return [(None, data.decode("utf-8", errors="replace"))]


def _parse_pdf(path: Path) -> list:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParseError("PDF parola korumalı")
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:       # tek bir bozuk sayfa tum belgeyi dusurmesin
                text = ""
            pages.append((i, text))
        return pages
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"PDF okunamadı: {type(e).__name__}") from e


def _parse_docx(path: Path) -> list:
    try:
        with zipfile.ZipFile(path) as z:      # DOCX bir zip'tir: acilmis toplam boyutu ONCE denetle
            if sum(i.file_size for i in z.infolist()) > config.MAX_DOCX_UNCOMPRESSED_MB * 1024 * 1024:
                raise ParseError("DOCX açıldığında çok büyük (sıkıştırma bombası olabilir)")
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:                       # tablo metinleri paragraflardan sonra eklenir
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [(None, "\n\n".join(parts))]
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"DOCX okunamadı: {type(e).__name__}") from e
